"""Saved players & groups, /database and the Excel export."""

from io import BytesIO

import discord
from openpyxl import load_workbook

from bot.export import build_workbook

from .test_bot import FakeChannel, FakeGuild, FakeInteraction, bot, scan_result, use_fake_roblox  # noqa: F401

GUILD = 1


def record(bot, result, **kw):  # noqa: F811
    from bot.db import GuildConfig

    defaults = {"depth": "full", "found_via": "/check", "checked_by": 5}
    return result.record(GuildConfig(guild_id=GUILD), **{**defaults, **kw})


async def get_player(bot, roblox_id):  # noqa: F811
    rows = await bot.db.list_players(GUILD, search=str(roblox_id))
    return rows[0] if rows else None


async def test_quick_checks_never_overwrite_full_checks(bot):  # noqa: F811
    alt = scan_result(age_days=2, badges=0, friends=0, groups=0, description="")
    alt.account.user_id = 77
    await bot.db.save_players(GUILD, [record(bot, alt)])
    first = await get_player(bot, 77)
    assert first["depth"] == "full" and first["flags"] == "alt" and first["friends"] == 0 and first["times_seen"] == 1

    quick = scan_result(age_days=2, badges=0, friends=None, groups=None, description="new bio")
    quick.account.user_id = 77
    await bot.db.save_players(GUILD, [record(bot, quick, depth="quick", found_via="friend of @Someone", checked_by=None)])
    after = await get_player(bot, 77)
    assert after["depth"] == "full" and after["friends"] == 0  # full results kept
    assert after["alt_score"] == first["alt_score"]
    assert after["bio"] == "new bio" and after["times_seen"] == 2  # but the profile itself is fresh
    assert after["found_via"] == "/check" and after["last_checked_by"] == 5

    # A later full check does replace an earlier quick one.
    fresh = scan_result(age_days=900, badges=150)
    fresh.account.user_id = 88
    await bot.db.save_players(GUILD, [record(bot, fresh, depth="quick", found_via="friend of @X")])
    await bot.db.save_players(GUILD, [record(bot, fresh)])
    assert (await get_player(bot, 88))["depth"] == "full"


async def test_player_filters_and_search(bot):  # noqa: F811
    nsfw = scan_result(description="add me 4 fun 18+ horny")
    nsfw.account.user_id, nsfw.account.username = 1, "Freaky_One"
    nsfw.account.display_name = "Freaky"
    alt = scan_result(age_days=2, badges=0, friends=0, groups=0, description="")
    alt.account.user_id, alt.account.username = 2, "Player99999"
    alt.account.display_name = "Player"
    fine = scan_result()
    fine.account.user_id, fine.account.username = 3, "BuilderKid"
    await bot.db.save_players(GUILD, [record(bot, r) for r in (nsfw, alt, fine)])

    assert await bot.db.count_players(GUILD) == 3
    assert [r["roblox_id"] for r in await bot.db.list_players(GUILD, show="flagged")] in ([1, 2], [2, 1])
    assert [r["roblox_id"] for r in await bot.db.list_players(GUILD, show="nsfw")] == [1]
    assert [r["roblox_id"] for r in await bot.db.list_players(GUILD, show="alts")] == [2]
    assert [r["roblox_id"] for r in await bot.db.list_players(GUILD, search="@builder")] == [3]
    assert [r["roblox_id"] for r in await bot.db.list_players(GUILD, search="2")] == [2]  # by ID
    assert await bot.db.count_players(GUILD + 1) == 0  # other servers can't see it


async def test_groups_are_saved_and_rescans_update(bot):  # noqa: F811
    row = {"group_id": 555, "name": "Cool Builders", "member_count": 6000, "owner": "Boss", "nsfw_score": 0,
           "nsfw_matches": "", "members_checked": 50, "nsfw_members": 1, "alt_members": 3, "suspicious_members": 4,
           "scanned_by": 5}
    await bot.db.save_group(GUILD, row)
    await bot.db.save_group(GUILD, {**row, "members_checked": 100, "alt_members": 6})
    groups = await bot.db.list_groups(GUILD)
    assert len(groups) == 1 and groups[0]["times_scanned"] == 2 and groups[0]["alt_members"] == 6
    counts = await bot.db.database_counts(GUILD)
    assert counts["groups"] == 1 and counts["players"] == 0


async def test_workbook(bot):  # noqa: F811
    evil = scan_result(description='=HYPERLINK("http://evil.example","click")\x07 hi')
    evil.account.user_id, evil.account.username = 9, "EvilBio"
    await bot.db.save_players(GUILD, [record(bot, evil)])
    await bot.db.save_group(GUILD, {"group_id": 555, "name": "-cmd|' /C calc'!A0", "member_count": 1, "owner": None,
                                    "nsfw_score": 0, "nsfw_matches": "", "members_checked": 1, "nsfw_members": 0,
                                    "alt_members": 0, "suspicious_members": 0, "scanned_by": 5})
    await bot.db.open_case(GUILD, 9, "EvilBio", ["manual"], 10, 0, {}, 5)
    tables = await bot.db.export_tables(GUILD)

    data = build_workbook("My Server", tables, lambda uid: f"Mod {uid}" if uid else "")
    book = load_workbook(BytesIO(data))
    assert book.sheetnames == ["Summary", "Players", "Groups", "Cases", "Case notes", "Tracked players", "Trusted",
                               "Activity log", "Filter words"]
    players = book["Players"]
    headers = [c.value for c in players[1]]
    row = {h: c for h, c in zip(headers, players[2], strict=True)}
    assert row["Username"].value == "EvilBio"
    assert row["Bio"].data_type == "s" and row["Bio"].value.startswith("=HYPERLINK")  # text, not a formula
    assert "\x07" not in row["Bio"].value
    assert row["Profile"].hyperlink.target == "https://www.roblox.com/users/9/profile"
    assert row["Last checked by"].value == "Mod 5"
    assert row["Last checked (UTC)"].is_date
    assert book["Groups"]["B2"].data_type == "s"
    assert book["Cases"]["I2"].value == "Mod 5"
    assert book["Summary"]["B1"].value == "My Server"


class ExportGuild(FakeGuild):
    name = "My Server"

    def get_member(self, user_id):
        return type("M", (), {"name": "alex", "display_name": "Alex"})() if user_id == 5 else None

    async def fetch_member(self, user_id):
        raise discord.NotFound(type("R", (), {"status": 404, "reason": "Not Found"})(), "Unknown Member")


async def test_database_commands(bot):  # noqa: F811
    cog = bot.get_cog("Database")
    guild = ExportGuild(FakeChannel())
    fine = scan_result()
    await bot.db.save_players(GUILD, [record(bot, fine)])

    summary = FakeInteraction(guild, type("U", (), {"id": 5})())
    await cog.summary.callback(cog, summary)
    embed = summary.response.sent[0][1]["embed"]
    assert "**1** saved" in embed.fields[0].value

    listing = FakeInteraction(guild, type("U", (), {"id": 5})())
    await cog.players.callback(cog, listing, "all", None, 5)
    assert "Page 1/1" in listing.response.sent[0][1]["embed"].footer.text

    export = FakeInteraction(guild, type("U", (), {"id": 5})())
    await cog.export.callback(cog, export)
    content, kwargs = export.followup.sent[0]
    assert kwargs["file"].filename.startswith("robloxchecker-My-Server-") and kwargs["file"].filename.endswith(".xlsx")
    book = load_workbook(kwargs["file"].fp)
    assert book["Players"]["W2"].value == "Alex (@alex)"
    assert "1 players" in content


async def test_checks_save_players(bot):  # noqa: F811
    from datetime import timedelta

    from bot.models import utcnow

    from .test_groupscan import FakeRobloxGroups
    from .test_roblox import USER, FakeRoblox, badge_pages, routes

    use_fake_roblox(bot, FakeRoblox(routes(**{
        "badges.roblox.com/v1/users/42/badges": badge_pages(0),
        "users.roblox.com/v1/users/42": {**USER, "created": (utcnow() - timedelta(days=2)).isoformat()},
    })))
    check = bot.get_cog("Check")
    interaction = FakeInteraction(FakeGuild(FakeChannel()), type("U", (), {"id": 5})())
    await check.check.callback(check, interaction, "SneakyAlt")
    saved = await get_player(bot, 42)
    assert saved["found_via"] == "/check" and saved["depth"] == "full" and saved["last_checked_by"] == 5

    use_fake_roblox(bot, FakeRobloxGroups())
    from .test_groupscan import ScanInteraction

    groups = bot.get_cog("Groups")
    await groups.groupscan.callback(groups, ScanInteraction(FakeGuild(FakeChannel()), type("U", (), {"id": 6})()), "555", 50, False)
    saved_group = (await bot.db.list_groups(GUILD))[0]
    assert saved_group["name"] == "Cool Builders" and saved_group["members_checked"] == 6
    assert (await get_player(bot, 102))["found_via"] == "group scan: Cool Builders"
