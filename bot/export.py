"""Everything a server has saved, as one Excel workbook (one tab per kind of data)."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

Name = Callable[[int | None], str]  # Discord user ID -> readable name

DATE_FORMAT = "yyyy-mm-dd hh:mm"
HEADER_FILL = PatternFill("solid", fgColor="2B4FD0")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="2B4FD0", underline="single")


class Link(str):
    """A cell value that should become a clickable hyperlink."""


def _when(value: str | None) -> datetime | None:
    if not value:
        return None
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo:
        parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
    return parsed  # Excel has no time zones; every date column says (UTC)


def _yes(value: Any) -> str:
    return "yes" if value else "no"


def _profile(roblox_id: int) -> Link:
    return Link(f"https://www.roblox.com/users/{roblox_id}/profile")


def _flags(value: str) -> str:
    names = {"alt": "Likely alt", "nsfw": "NSFW", "manual": "Manual"}
    return ", ".join(names.get(f, f) for f in value.split(",") if f)


def _badges(row) -> str | int:
    if row["badges"] is None:
        return "private"
    return f"{row['badges']}+" if row["badges_more"] else row["badges"]


def sheets(tables: dict[str, list], name: Name) -> list[tuple[str, list[str], list[list[Any]]]]:
    """``[(tab name, headers, rows)]`` built from :meth:`Database.export_tables`."""
    return [
        ("Players", [
            "Roblox ID", "Username", "Display name", "Profile", "Flags", "Alt score", "Alt verdict", "NSFW score",
            "NSFW matches", "Account created (UTC)", "Badges", "Friends", "Followers", "Groups", "Banned", "Verified",
            "Bio", "Found via", "Check type", "First seen (UTC)", "Last checked (UTC)", "Times seen", "Last checked by",
        ], [[
            r["roblox_id"], r["username"], r["display_name"], _profile(r["roblox_id"]), _flags(r["flags"]),
            r["alt_score"], r["alt_verdict"], r["nsfw_score"], r["nsfw_matches"], _when(r["account_created"]),
            _badges(r), r["friends"], r["followers"], r["groups_count"], _yes(r["is_banned"]), _yes(r["verified"]),
            r["bio"], r["found_via"], "Full" if r["depth"] == "full" else "Quick (friend)", _when(r["first_seen"]),
            _when(r["last_checked"]), r["times_seen"], name(r["last_checked_by"]),
        ] for r in tables["players"]]),
        ("Groups", [
            "Group ID", "Name", "Link", "Members", "Owner", "Group name/description NSFW score", "NSFW matches",
            "Members checked", "NSFW members", "Likely alts", "Suspicious", "First scanned (UTC)", "Last scanned (UTC)",
            "Times scanned", "Last scanned by",
        ], [[
            r["group_id"], r["name"], Link(f"https://www.roblox.com/communities/{r['group_id']}"), r["member_count"],
            r["owner"], r["nsfw_score"], r["nsfw_matches"], r["members_checked"], r["nsfw_members"], r["alt_members"],
            r["suspicious_members"], _when(r["first_scanned"]), _when(r["last_scanned"]), r["times_scanned"],
            name(r["last_scanned_by"]),
        ] for r in tables["groups"]]),
        ("Cases", [
            "Case #", "Username", "Roblox ID", "Profile", "Reasons", "Status", "Alt score", "NSFW score",
            "Opened by", "Opened (UTC)", "Claimed by", "Resolved by", "Resolved (UTC)", "Notes",
        ], [[
            r["id"], r["roblox_name"], r["roblox_id"], _profile(r["roblox_id"]), _flags(r["reasons"]), r["status"],
            r["alt_score"], r["nsfw_score"], name(r["opened_by"]) if r["opened_by"] else "Auto-flag",
            _when(r["opened_at"]), name(r["claimed_by"]) if r["claimed_by"] else "", name(r["resolved_by"]) if r["resolved_by"] else "",
            _when(r["resolved_at"]), r["note_count"],
        ] for r in tables["cases"]]),
        ("Case notes", ["Case #", "Username", "Written by", "Note", "Written (UTC)"], [[
            r["case_id"], r["roblox_name"], name(r["author_id"]), r["body"], _when(r["created_at"]),
        ] for r in tables["notes"]]),
        ("Tracked players", [
            "Username", "Roblox ID", "Profile", "Why", "Added by", "Added (UTC)", "Last checked (UTC)", "Badges",
            "Alt score", "NSFW score",
        ], [[
            r["roblox_name"], r["roblox_id"], _profile(r["roblox_id"]), r["note"], name(r["added_by"]), _when(r["added_at"]),
            _when(r["last_checked_at"]), "private" if r["last_badge_count"] == -1 else r["last_badge_count"],
            r["last_alt_score"], r["last_nsfw_score"],
        ] for r in tables["watchlist"]]),
        ("Trusted", ["Username", "Roblox ID", "Profile", "Why", "Added by", "Added (UTC)"], [[
            r["roblox_name"], r["roblox_id"], _profile(r["roblox_id"]), r["reason"], name(r["added_by"]), _when(r["added_at"]),
        ] for r in tables["trusted"]]),
        ("Activity log", ["When (UTC)", "Who", "Action", "What", "Details"], [[
            _when(r["created_at"]), name(r["actor_id"]) if r["actor_id"] else "Bot", r["action"], r["target"], r["details"],
        ] for r in tables["activity"]]),
        ("Filter words", ["Word", "Points", "Match", "Category", "Changed by", "Changed (UTC)"], [[
            r["term"], r["weight"] or "switched off", r["match"], r["category"], name(r["added_by"]), _when(r["added_at"]),
        ] for r in tables["terms"]]),
    ]


def _write(cell, value: Any) -> None:
    if isinstance(value, Link):
        cell.value = str(value)
        cell.hyperlink = str(value)
        cell.font = LINK_FONT
        return
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32000]
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # text a player wrote (like "=HYPERLINK(...)") must never run as a formula
    elif isinstance(value, datetime):
        cell.number_format = DATE_FORMAT


def build_workbook(server_name: str, tables: dict[str, list], name: Name, *, note: str = "") -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    built = sheets(tables, name)
    rows = [
        ["RobloxChecker export", server_name],
        ["Exported (UTC)", datetime.now(timezone.utc).replace(tzinfo=None)],
        *([["Note", note]] if note else []),
        [],
        ["Tab", "Rows"],
        *[[title, len(data)] for title, _, data in built],
    ]
    for r, values in enumerate(rows, start=1):
        for c, value in enumerate(values, start=1):
            _write(summary.cell(row=r, column=c), value)
    summary["A1"].font = Font(bold=True, size=14)
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 40

    for title, headers, data in built:
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL
            cell.alignment = Alignment(vertical="center")
        for r, values in enumerate(data, start=2):
            for c, value in enumerate(values, start=1):
                _write(sheet.cell(row=r, column=c), value)
        sheet.freeze_panes = "A2"
        if data:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(data) + 1}"
        for c, header in enumerate(headers, start=1):
            longest = max([len(header)] + [len(str(row[c - 1])) for row in data[:500] if row[c - 1] is not None])
            sheet.column_dimensions[get_column_letter(c)].width = min(50, max(10, longest + 2))

    out = BytesIO()
    workbook.save(out)
    return out.getvalue()
